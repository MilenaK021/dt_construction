"""
tender_mailer.py  ->  api/tender_mailer.py
Sends tender details email with a Word (.docx) attachment.
"""

import os, io, smtplib
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text      import MIMEText
from email.mime.base      import MIMEBase
from email                import encoders

from fastapi   import APIRouter, HTTPException
from pydantic  import BaseModel
from dotenv    import load_dotenv

from docx                  import Document
from docx.shared           import Pt, RGBColor, Cm
from docx.oxml.ns          import qn
from docx.oxml             import OxmlElement

load_dotenv()
router = APIRouter()


# ── Pydantic models ────────────────────────────────────────────────────────────

class TenderData(BaseModel):
    id:           str
    title:        str
    customer:     str
    region:       str
    budget:       str
    deadline:     str
    category:     str
    description:  str
    requirements: list[str] = []
    published:    str = ""

class TenderEmailRequest(BaseModel):
    tender:    TenderData
    recipient: str


# ── Helpers ────────────────────────────────────────────────────────────────────

def _set_color(run, hex_color: str):
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    run.font.color.rgb = RGBColor(r, g, b)


def _add_separator(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    pPr    = p._p.get_or_add_pPr()
    pBdr   = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "e4e8ef")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_heading(doc, text, size=13, bold=True, color="1a2535",
                 space_before=6, space_after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.bold      = bold
    run.font.size = Pt(size)
    _set_color(run, color)


def _add_para(doc, text, size=10.5, color="374151", space_after=4, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold      = bold
    _set_color(run, color)


# ── Docx builder ───────────────────────────────────────────────────────────────

def _build_docx(tender: TenderData) -> bytes:
    doc = Document()

    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # Header line
    h = doc.add_paragraph()
    h.paragraph_format.space_after = Pt(2)
    r1 = h.add_run("DIGITAL TWIN  ")
    r1.bold = True
    r1.font.size = Pt(11)
    _set_color(r1, "2563eb")
    r2 = h.add_run(
        "Construction Management  ·  " +
        datetime.now().strftime("%d.%m.%Y %H:%M")
    )
    r2.font.size = Pt(9)
    _set_color(r2, "9ca3af")
    _add_separator(doc)

    # Category + Title
    doc.add_paragraph()
    cat_p = doc.add_paragraph()
    cat_p.paragraph_format.space_after = Pt(4)
    cat_r = cat_p.add_run(tender.category.upper())
    cat_r.font.size = Pt(9)
    cat_r.bold = True
    _set_color(cat_r, "2563eb")

    _add_heading(doc, tender.title, size=15, color="1a2535",
                 space_before=2, space_after=4)

    id_p = doc.add_paragraph()
    id_p.paragraph_format.space_after = Pt(10)
    id_r = id_p.add_run(
        "Тендер #" + tender.id +
        "  ·  Опубликован: " + (tender.published or "—")
    )
    id_r.font.size = Pt(9)
    _set_color(id_r, "8896ab")
    _add_separator(doc)

    # Key details table
    doc.add_paragraph()
    _add_heading(doc, "Основные параметры", size=10, color="2563eb",
                 space_before=4, space_after=6)

    rows_data = [
        ("Заказчик",    tender.customer,           "1a2535", False),
        ("Регион",      tender.region,             "1a2535", False),
        ("Бюджет",      tender.budget,             "16a34a", True),
        ("Срок подачи", tender.deadline,           "dc2626", True),
        ("Опубликован", tender.published or "—",  "1a2535", False),
        ("Категория",   tender.category,           "1a2535", False),
    ]
    table = doc.add_table(rows=len(rows_data), cols=2)
    table.style = "Table Grid"

    for i, (label, val, val_color, val_bold) in enumerate(rows_data):
        row = table.rows[i]

        lc = row.cells[0]
        lc.width = Cm(4)
        lr = lc.paragraphs[0].add_run(label)
        lr.bold = True
        lr.font.size = Pt(9)
        _set_color(lr, "6b7a8f")

        vc = row.cells[1]
        vr = vc.paragraphs[0].add_run(val)
        vr.font.size = Pt(10)
        vr.bold = val_bold
        _set_color(vr, val_color)

        if i % 2 == 0:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                shd  = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "f8f9fc")
                shd.set(qn("w:val"),  "clear")
                tcPr.append(shd)

    # Description
    doc.add_paragraph()
    _add_heading(doc, "Описание", size=10, color="2563eb",
                 space_before=10, space_after=4)
    _add_para(doc, tender.description, size=10.5)

    # Requirements
    if tender.requirements:
        _add_heading(doc, "Требования к участнику", size=10, color="2563eb",
                     space_before=8, space_after=4)
        for req in tender.requirements:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(req)
            run.font.size = Pt(10)
            _set_color(run, "374151")

    # Footer
    doc.add_paragraph()
    _add_separator(doc)
    foot = doc.add_paragraph()
    foot.paragraph_format.space_before = Pt(4)
    fr = foot.add_run(
        "Документ сформирован автоматически · Digital Twin Construction Management · "
        "Источник: goszakup.gov.kz (mock data)"
    )
    fr.font.size = Pt(8)
    _set_color(fr, "9ca3af")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ── HTML email body ────────────────────────────────────────────────────────────

def _detail_rows(tender: TenderData) -> str:
    rows = [
        ("Заказчик",    tender.customer),
        ("Регион",      tender.region),
        ("Бюджет",
         '<span style="color:#16a34a;font-weight:700">' + tender.budget + '</span>'),
        ("Срок подачи",
         '<span style="color:#dc2626;font-weight:600">' + tender.deadline + '</span>'),
    ]
    parts = []
    for label, val in rows:
        parts.append(
            '<tr style="border-bottom:1px solid #e4e8ef">' +
            '<td style="padding:11px 16px;font-size:11px;font-weight:600;' +
            'color:#8896ab;width:130px;white-space:nowrap">' + label + '</td>' +
            '<td style="padding:11px 16px;font-size:13px;color:#1a2535;' +
            'font-weight:500">' + val + '</td></tr>'
        )
    return "".join(parts)


def _reqs_block(tender: TenderData, reqs_html: str) -> str:
    if not tender.requirements:
        return ""
    return (
        '<tr><td style="padding:0 32px 24px">' +
        '<div style="font-size:12px;font-weight:600;color:#2563eb;' +
        'text-transform:uppercase;letter-spacing:.06em;margin-bottom:10px">' +
        'Требования</div>' +
        '<ul style="margin:0;padding-left:18px;font-size:13px;' +
        'color:#374151;line-height:1.8">' + reqs_html + '</ul>' +
        '</td></tr>'
    )


def _html_email(tender: TenderData) -> str:
    reqs_html = "".join("<li>" + r + "</li>" for r in tender.requirements)
    return (
        '''<!DOCTYPE html><html lang="ru">
<head><meta charset="UTF-8"></head>
<body style="margin:0;padding:0;background:#f0f2f5;font-family:Arial,sans-serif">
<table width="100%" cellpadding="0" cellspacing="0" style="background:#f0f2f5;padding:32px 0">
<tr><td align="center">
<table width="600" cellpadding="0" cellspacing="0"
  style="background:#fff;border-radius:16px;overflow:hidden;border:1px solid #e4e8ef">

<tr><td style="background:#0f1e30;padding:24px 32px">
  <table width="100%" cellpadding="0" cellspacing="0"><tr>
    <td><span style="color:#60a5fa;font-size:13px;font-weight:600">DIGITAL TWIN</span>
        <div style="color:rgba(255,255,255,.5);font-size:11px;margin-top:2px">Construction Management</div></td>
    <td align="right"><span style="background:rgba(59,130,246,.2);color:#60a5fa;
      font-size:11px;font-weight:600;padding:4px 12px;border-radius:20px;
      border:1px solid rgba(59,130,246,.3)">Новый тендер</span></td>
  </tr></table>
</td></tr>

<tr><td style="padding:28px 32px 20px">
  <div style="font-size:11px;color:#6b7a8f;font-weight:600;text-transform:uppercase;
    letter-spacing:.07em;margin-bottom:8px">''' + tender.category + '''</div>
  <h1 style="margin:0 0 8px;font-size:20px;font-weight:700;color:#1a2535;line-height:1.3">''' +
  tender.title + '''</h1>
  <div style="color:#8896ab;font-size:13px">Тендер #''' + tender.id +
  '''  ·  Опубликован: ''' + (tender.published or "—") + '''</div>
</td></tr>

<tr><td style="padding:0 32px 24px">
  <table width="100%" cellpadding="0" cellspacing="0"
    style="background:#f8f9fc;border:1px solid #e4e8ef;border-radius:12px;overflow:hidden">
    ''' + _detail_rows(tender) + '''
  </table>
</td></tr>

<tr><td style="padding:0 32px 24px">
  <div style="font-size:12px;font-weight:600;color:#2563eb;text-transform:uppercase;
    letter-spacing:.06em;margin-bottom:10px">Описание</div>
  <div style="font-size:13px;color:#374151;line-height:1.7">''' + tender.description + '''</div>
</td></tr>

''' + _reqs_block(tender, reqs_html) + '''

<tr><td style="padding:0 32px 32px">
  <div style="background:#eff6ff;border:1px solid #bfdbfe;border-radius:12px;padding:16px 20px">
    <div style="font-size:13px;color:#1d4ed8;font-weight:500;margin-bottom:4px">
      📎 Подробная документация прикреплена к письму (.docx)
    </div>
    <div style="font-size:12px;color:#6b7a8f">
      Word-файл содержит полные условия тендера.
    </div>
  </div>
</td></tr>

<tr><td style="background:#f8f9fc;border-top:1px solid #e4e8ef;padding:16px 32px">
  <div style="font-size:11px;color:#9ca3af;line-height:1.6">
    Сообщение сформировано автоматически · Digital Twin Construction Management<br>
    Источник данных: goszakup.gov.kz
  </div>
</td></tr>

</table></td></tr></table>
</body></html>'''
    )


# ── SMTP sender ────────────────────────────────────────────────────────────────

def _send_tender_email(recipient: str, tender: TenderData, docx_bytes: bytes):
    host = os.getenv("SMTP_HOST", "smtp.gmail.com")
    port = int(os.getenv("SMTP_PORT", 587))
    user = os.getenv("SMTP_USER", "")
    pwd  = os.getenv("SMTP_PASS", "")
    frm  = os.getenv("SMTP_FROM", user)

    if not user or not pwd:
        raise ValueError("SMTP_USER and SMTP_PASS must be set in .env")

    msg = MIMEMultipart("mixed")
    msg["Subject"] = "Тендер: " + tender.title + " — детали и документация"
    msg["From"]    = frm
    msg["To"]      = recipient

    alt = MIMEMultipart("alternative")
    alt.attach(MIMEText(_html_email(tender), "html", "utf-8"))
    msg.attach(alt)

    safe_name = tender.title[:40].replace(" ", "_").replace("/", "-")
    att = MIMEBase(
        "application",
        "vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    att.set_payload(docx_bytes)
    encoders.encode_base64(att)
    att.add_header(
        "Content-Disposition", "attachment",
        filename="Tender_" + tender.id + "_" + safe_name + ".docx"
    )
    msg.attach(att)

    with smtplib.SMTP(host, port) as smtp:
        smtp.ehlo()
        smtp.starttls()
        smtp.login(user, pwd)
        smtp.sendmail(frm, [recipient], msg.as_string())


# ── Route ──────────────────────────────────────────────────────────────────────

@router.post("/tenders/send-email")
def send_tender_email(req: TenderEmailRequest):
    if not req.recipient or "@" not in req.recipient:
        raise HTTPException(status_code=400,
            detail="Укажите email директора в профиле (раздел «Профиль директора»)")
    try:
        docx_bytes = _build_docx(req.tender)
        _send_tender_email(req.recipient, req.tender, docx_bytes)
    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except smtplib.SMTPAuthenticationError:
        raise HTTPException(status_code=500,
            detail="SMTP authentication failed. Проверьте SMTP_USER / SMTP_PASS в .env")
    except Exception as e:
        raise HTTPException(status_code=500, detail="Email error: " + str(e))

    return {"status": "sent", "recipient": req.recipient, "tender_id": req.tender.id}